"""
Assemble Phase I investigation report as .docx.
Sections: cover, event description, phase I narrative,
          escalation recommendation, signature block, disclaimer.
No CAPA. No Phase II content.
"""
from __future__ import annotations
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


class ReportBuilder:
    def __init__(self, config: dict):
        self.config = config

    def build(
        self,
        event: dict,
        sections: dict,
        out_dir: Path,
        shap_png_path: str | None = None,
    ) -> Path:
        doc = Document()
        self._style(doc)
        self._header_footer(doc, event)

        self._cover(doc, event)
        self._event_description(doc, event)
        self._phase1(doc, sections["phase1"])
        if shap_png_path and Path(shap_png_path).exists():
            self._shap_figure(doc, shap_png_path)
        self._escalation(doc, sections["escalation"])
        self._signatures(doc)
        self._disclaimer(doc)

        fname = (
            f"{event.get('event_type','EVT')}_"
            f"{event.get('batch_number','BATCH')}_"
            f"{event.get('event_id','ID')}_PhaseI.docx"
        )
        path = out_dir / fname
        doc.save(str(path))
        return path

    # ── private helpers ───────────────────────────────────────────────────────

    def _style(self, doc):
        style = doc.styles["Normal"]
        style.font.name = self.config["report"].get("font", "Calibri")
        style.font.size = Pt(11)

    def _header_footer(self, doc, event):
        sec = doc.sections[0]
        hp = sec.header.paragraphs[0]
        hp.clear()
        run = hp.add_run(self.config["report"].get("watermark_text", "AI-ASSISTED DRAFT"))
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xCC, 0x00, 0x00)
        run.bold = True

        fp = sec.footer.paragraphs[0]
        fp.clear()
        fp.add_run(
            f"Event: {event.get('event_id','')}  |  "
            f"Phase I Investigation Report  |  "
            f"Model: {self.config['llm']['model']}  |  CONFIDENTIAL"
        ).font.size = Pt(8)

    def _cover(self, doc, event):
        doc.add_paragraph()
        t = doc.add_paragraph()
        t.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = t.add_run("PHASE I LABORATORY INVESTIGATION REPORT")
        r.bold = True; r.font.size = Pt(16)

        doc.add_paragraph()
        tbl = doc.add_table(rows=8, cols=2)
        tbl.style = "Table Grid"
        rows = [
            ("Event ID",         event.get("event_id", "")),
            ("Event Type",       event.get("event_type", "")),
            ("Batch Number",     event.get("batch_number", "")),
            ("Product Code",     event.get("product_code", "")),
            ("Parameter",        event.get("parameter", "")),
            ("Result / Spec",    f"{event.get('result','')} {event.get('unit','')} / {event.get('spec_limit','')}"),
            ("Triggered At",     str(event.get("triggered_at",""))[:19]),
            ("Report Generated", datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC")),
        ]
        for i, (lbl, val) in enumerate(rows):
            tbl.rows[i].cells[0].text = lbl
            tbl.rows[i].cells[1].text = str(val)
            tbl.rows[i].cells[0].paragraphs[0].runs[0].bold = True
        doc.add_page_break()

    def _event_description(self, doc, event):
        doc.add_heading("1. Event Description", level=1)
        p = doc.add_paragraph()
        p.add_run(
            f"On {str(event.get('triggered_at',''))[:10]}, a "
            f"{event.get('event_type')} event was recorded for batch "
            f"{event.get('batch_number')} ({event.get('product_code')}). "
            f"The parameter “{event.get('parameter')}” yielded "
            f"{event.get('result')} {event.get('unit')} against a specification "
            f"of {event.get('spec_limit')}."
        )
        if event.get("deviation_pct"):
            p.add_run(
                f" This represents a {event['deviation_pct']}% deviation from "
                f"the specification midpoint ({event.get('deviation_direction','')} deviation)."
            )

        if event.get("shap_top_features"):
            doc.add_heading("SHAP Anomaly Evidence (ML-derived)", level=2)
            tbl = doc.add_table(rows=1 + len(event["shap_top_features"]), cols=3)
            tbl.style = "Table Grid"
            for j, h in enumerate(["Feature", "SHAP Value", "Feature Value"]):
                tbl.rows[0].cells[j].text = h
                tbl.rows[0].cells[j].paragraphs[0].runs[0].bold = True
            for i, feat in enumerate(event["shap_top_features"], 1):
                tbl.rows[i].cells[0].text = feat.get("feature", "")
                tbl.rows[i].cells[1].text = f"{feat.get('shap_value', 0):.4f}"
                tbl.rows[i].cells[2].text = f"{feat.get('feature_value', 0):.4f}"

    def _phase1(self, doc, text: str):
        doc.add_heading("2. Phase I — Laboratory Investigation", level=1)
        for para in text.strip().split("

"):
            if para.strip():
                doc.add_paragraph(para.strip())

    def _shap_figure(self, doc, png_path: str):
        doc.add_heading("Figure 1: SHAP Feature Contribution (Force Plot)", level=2)
        doc.add_picture(png_path, width=Inches(6.0))
        cap = doc.add_paragraph("Figure 1. SHAP force plot for this event. "
                                 "Red bars push toward OOS/OOT; blue bars push toward Normal.")
        cap.style = "Caption" if "Caption" in [s.name for s in doc.styles] else "Normal"

    def _escalation(self, doc, text: str):
        doc.add_heading("3. Escalation Recommendation", level=1)
        for para in text.strip().split("

"):
            if para.strip():
                doc.add_paragraph(para.strip())
        doc.add_paragraph(
            "Note: Phase II manufacturing investigation and CAPA generation are outside "
            "the scope of this automated system. If escalation is recommended above, "
            "the responsible site quality team must initiate Phase II per site SOP."
        ).runs[0].italic = True

    def _signatures(self, doc):
        doc.add_heading("4. Review and Approval", level=1)
        doc.add_paragraph(
            "This report covers Phase I laboratory investigation only. "
            "QC and QA review is required before any batch disposition decision."
        )
        tbl = doc.add_table(rows=4, cols=3)
        tbl.style = "Table Grid"
        for j, h in enumerate(["Role", "Name (Print)", "Signature / Date"]):
            tbl.rows[0].cells[j].text = h
            tbl.rows[0].cells[j].paragraphs[0].runs[0].bold = True
        for i, role in enumerate(["Investigating Analyst", "QC Supervisor", "QA Reviewer"], 1):
            tbl.rows[i].cells[0].text = role

    def _disclaimer(self, doc):
        doc.add_paragraph()
        p = doc.add_paragraph()
        r = p.add_run(
            f"AI-ASSISTED DRAFT — PHASE I SCOPE ONLY. "
            f"Generated using {self.config['llm']['model']}. "
            "Content must be verified by a qualified pharmaceutical professional. "
            "This document does not constitute a batch release decision."
        )
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        r.italic = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
